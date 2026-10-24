"""Excel 值班表导出。

模板约定（已核对实际文件）：

- ``A2`` 标记所在格 = 该部门可排的班次格。周一~周六各 6 段节次共 36 格，
  周日只有 3 格（1-2、3-4、5-6），合计 39 格。
- **校验只认周一~周六这 36 格**：周日永远不排班（见
  :func:`app.scheduler._eligible_slots`），所以模板里有没有周日格都不影响
  导出；新模板把周日列删掉也不该报"模板有问题"。
- 值班人姓名写在 ``A2`` 格右侧一列（D→E、F→G … N→O、P→Q）。
- 模板里本身带着旧名字（``O5``、``E17`` …），导出前必须全部清空，
  否则新排班会被旧数据污染——这正是旧脚本排错班的原因之一。

``A3`` 那几行属于另一个部门/另一套标记，本模块不碰。
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Dict, Mapping, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from app.jiaowu import DAYS, SLOTS, _band_of, detect_slot, parse_row_period

# 节次 -> 模板行号
SLOT_ROW: Dict[str, int] = {
    "1-2": 5,
    "3-4": 7,
    "5-6": 11,
    "7-8": 13,
    "9-10": 15,
    "11-12": 17,
}
ROW_SLOT: Dict[int, str] = {row: slot for slot, row in SLOT_ROW.items()}

# 星期 -> A2 标记所在列
DAY_COL: Dict[str, str] = {
    "周一": "D",
    "周二": "F",
    "周三": "H",
    "周四": "J",
    "周五": "L",
    "周六": "N",
    "周日": "P",
}

EXPECTED_CELL_COUNT = 39

# 校验口径：只要求周一~周六齐（各 6 段 = 36 格）。周日不排班，
# 模板里那 3 个周日格有就收下、没有也不报错。
REQUIRED_DAYS: Tuple[str, ...] = tuple(d for d in DAYS if d != "周日")
REQUIRED_CELL_COUNT = len(REQUIRED_DAYS) * len(SLOTS)  # 36

TIME_LABEL = "排班时间"
PLACE_LABEL = "值班地点"

Slot = Tuple[str, str]


class TemplateChanged(RuntimeError):
    """模板结构与预期不符：宁可报错，不要静默导出一张错表。"""


@dataclass
class TemplateLayout:
    markers: Dict[Slot, str]      # (星期, 节次) -> A2 标记格坐标
    name_cells: Dict[Slot, str]   # (星期, 节次) -> 姓名格坐标

    def __len__(self) -> int:
        return len(self.markers)

    @property
    def slots(self) -> set[Slot]:
        return set(self.markers)


MARKER = "A2"
DAY_ALIASES = {
    "一": "周一", "二": "周二", "三": "周三", "四": "周四",
    "五": "周五", "六": "周六", "日": "周日", "天": "周日",
}


def _marker_and_name(day: str, slot: str) -> tuple[str, str]:
    col = DAY_COL[day]
    row = SLOT_ROW[slot]
    marker = f"{col}{row}"
    name = f"{get_column_letter(openpyxl.utils.column_index_from_string(col) + 1)}{row}"
    return marker, name


def _right(cell_coord: str) -> str:
    """标记格右侧一列的坐标（姓名写在这里）。"""
    column = "".join(ch for ch in cell_coord if ch.isalpha())
    row = "".join(ch for ch in cell_coord if ch.isdigit())
    return f"{get_column_letter(openpyxl.utils.column_index_from_string(column) + 1)}{row}"


def _detect_layout(sheet) -> Dict[Slot, str]:
    """按内容自动识别布局：不依赖固定行列号，只要表头写星期、左边写节次即可。

    换模板时（内容变了、排版没变，或者整体挪了一两行）靠这个救回来，
    而不是让用户对着"找到 0 个 A2 可填格"发呆。
    """
    # 1) 表头：列号 -> 星期
    day_of_col: Dict[int, str] = {}
    for row in sheet.iter_rows():
        for cell in row:
            m = re.search(r"(?:星期|周)\s*([一二三四五六日天])", str(cell.value or ""))
            if m:
                day = DAY_ALIASES.get(m.group(1))
                if day and cell.column not in day_of_col:
                    day_of_col[cell.column] = day
    if not day_of_col:
        return {}

    # 2) A2 标记格 + 同行左侧的节次标签
    found: Dict[Slot, str] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if str(cell.value or "").strip() != MARKER:
                continue
            day = day_of_col.get(cell.column)
            if day is None:
                # 标记格偶尔比表头偏一列，就近找一次
                for delta in (-1, 1, -2, 2):
                    if day_of_col.get(cell.column + delta):
                        day = day_of_col[cell.column + delta]
                        break
            if day is None:
                continue
            slot = None
            for left in range(cell.column - 1, 0, -1):
                text = str(sheet.cell(row=cell.row, column=left).value or "")
                label = detect_slot(text)
                if not label:
                    period = parse_row_period(text)
                    label = _band_of(period) if period else None
                if label:
                    slot = label
                    break
            if slot is None:
                continue
            found.setdefault((day, slot), cell.coordinate)
    return found


def _required_missing(found: Mapping[Slot, str]) -> list[Slot]:
    """周一~周六里还缺哪些格。周日不参与校验（永不排班）。"""
    return [
        (day, slot)
        for day in REQUIRED_DAYS
        for slot in SLOTS
        if (day, slot) not in found
    ]


def read_layout(template_path: Path | str) -> TemplateLayout:
    """扫描模板，返回可填格与姓名格坐标。结构不符时抛 TemplateChanged。

    先按标准坐标读（绝大多数情况，一次就中）；对不上再按内容自动识别，
    这样"只换表格内容、不动排版"的新模板可以直接用。

    周日格不是必需的：模板删掉周日列照样能用，只要周一~周六 36 格齐。
    """
    workbook = openpyxl.load_workbook(str(template_path))
    try:
        sheet = workbook.active
        found: Dict[Slot, str] = {}
        for day in DAYS:
            for slot in SLOTS:
                marker, _name = _marker_and_name(day, slot)
                if str(sheet[marker].value or "").strip() == MARKER:
                    found[(day, slot)] = marker
        if _required_missing(found):
            # 标准坐标不完整 → 按内容重新识别（换过模板 / 整体挪行）
            detected = _detect_layout(sheet)
            if len(detected) > len(found):
                found = detected
    finally:
        workbook.close()

    missing = _required_missing(found)
    if missing:
        raise TemplateChanged(
            f"模板里只找到 {len(found)} 个「{MARKER}」可填格，"
            f"周一~周六的 {REQUIRED_CELL_COUNT} 个必需格还缺 "
            f"{len(missing)} 个：{'、'.join(f'{day}{slot}' for day, slot in missing)}。"
            f"请确认表头有「星期一…星期六」、左边有「1-2节」这类节次、"
            f"可填格写着「{MARKER}」"
        )
    return TemplateLayout(
        markers=found,
        name_cells={key: _right(coord) for key, coord in found.items()},
    )


def schedule_filename(week: int, semester: str) -> str:
    """导出用的文件名，沿用原项目 ``build_schedule_filename`` 的命名规则。

    形如 ``勤工助学值班表 2026-2027 -1第7周.xlsx``：
    学期配置是 ``2026-2027-1`` 这样的三段式，拆成年份区间 + 学期号。
    任何一段缺失或不是 1/2 学期，就退回占位符（``20XX`` / ``X``），
    和原脚本的兜底行为一致，不会因为配置没填就出错。
    """
    parts = (semester or "").split("-")
    year1 = parts[0] if len(parts) >= 1 and parts[0] else "20XX"
    year2 = parts[1] if len(parts) >= 2 and parts[1] else "20XX"
    semester_num = parts[2] if len(parts) >= 3 else "X"
    week_str = str(week) if week > 0 else "X"
    semester_cn = semester_num if semester_num in ("1", "2") else "X"
    return f"勤工助学值班表 {year1}-{year2} -{semester_cn}第{week_str}周.xlsx"


def describe_template(template_path: Path | str) -> dict:
    """给界面看的模板摘要：文件名 + 识别到多少格。"""
    path = Path(template_path)
    info = {"path": str(path), "name": path.name, "exists": path.exists(),
            "cells": 0, "error": ""}
    if not path.exists():
        info["error"] = "模板文件不存在"
        return info
    try:
        info["cells"] = len(read_layout(path))
    except TemplateChanged as exc:
        info["error"] = str(exc)
    except Exception as exc:  # noqa: BLE001 - 损坏的文件不该让页面 500
        info["error"] = f"模板读不出来：{exc}"
    return info


def store_template(source: bytes, dest: Path) -> TemplateLayout:
    """校验通过才落盘：先写临时文件，读完布局没问题再替换。

    校验失败时现场一点副作用都不留，用户拿到的还是原来那份能用的模板。
    临时文件必须保留 ``.xlsx`` 后缀——openpyxl 按扩展名判断格式，改了就读不了。
    """
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    staging = dest.parent / f"{dest.stem}.uploading.xlsx"
    staging.write_bytes(source)
    try:
        layout = read_layout(staging)
    except Exception:
        staging.unlink(missing_ok=True)
        raise
    staging.replace(dest)
    return layout


def fillable_slots(template_path: Path | str) -> set[Slot]:
    return read_layout(template_path).slots


def export_schedule(
    template_path: Path | str,
    out_path: Path | str,
    assignments: Mapping[Slot, str],
    week: Optional[int] = None,
    today: Optional[date] = None,
) -> Path:
    """按模板导出值班表。``assignments`` 为 (星期, 节次) -> 姓名。"""
    layout = read_layout(template_path)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(str(template_path), str(out_path))

    workbook = openpyxl.load_workbook(str(out_path))
    sheet = workbook.active

    # 1) 先清空所有可填格的姓名，杜绝旧数据残留
    for slot_key, name_cell in layout.name_cells.items():
        sheet[name_cell] = None

    # 2) 再写入本周排班；不在可填格里的分配直接拒绝，避免写歪
    unknown = [key for key in assignments if key not in layout.markers]
    if unknown:
        workbook.close()
        out_path.unlink(missing_ok=True)
        raise ValueError(f"这些格子不在模板可填范围内：{unknown}")

    for (day, slot), name in assignments.items():
        sheet[layout.name_cells[(day, slot)]] = name

    if week is not None:
        stamp = today or date.today()
        sheet["A3"] = _compose_header(sheet["A3"].value, week, stamp)

    workbook.save(str(out_path))
    workbook.close()
    return out_path


def export_schedule_bytes(
    template_path: Path | str,
    assignments: Mapping[Slot, str],
    week: Optional[int] = None,
    today: Optional[date] = None,
    texts: Optional[Mapping[str, str]] = None,
) -> bytes:
    """按模板导出值班表为内存字节流，不落临时文件。

    与 ``export_schedule`` 同样的清空 + 校验逻辑，只是输出走 BytesIO。
    ``texts`` 是在线改过的文案（标题/地点/老师/备注），有值就覆盖模板原文。
    """
    layout = read_layout(template_path)
    workbook = openpyxl.load_workbook(str(template_path))
    sheet = workbook.active

    # 1) 先清空所有可填格的姓名，杜绝旧数据残留
    for slot_key, name_cell in layout.name_cells.items():
        sheet[name_cell] = None

    # 2) 再写入本周排班；不在可填格里的分配直接拒绝，避免写歪
    unknown = [key for key in assignments if key not in layout.markers]
    if unknown:
        workbook.close()
        raise ValueError(f"这些格子不在模板可填范围内：{unknown}")

    for (day, slot), name in assignments.items():
        sheet[layout.name_cells[(day, slot)]] = name

    if week is not None or texts:
        stamp_date = today or date.today()
        _apply_texts(sheet, texts or {}, week, stamp_date)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


# ------------------------------------------------------------ 可改文案（在线编辑）

#: 模板里允许在线改的文案字段（顺序即界面顺序）
TEXT_FIELD_LABELS: Dict[str, str] = {
    "title": "表格标题",
    "place": "值班地点",
    "teacher_name": "负责老师姓名",
    "teacher_phone": "负责老师联系电话",
    "note": "备注",
}

#: 单字段长度上限，防止有人把整篇 doc 贴进来
TEXT_FIELD_MAX_LEN = 600


def _merged_span(sheet, coord: str) -> Tuple[int, int, int, int]:
    for rng in sheet.merged_cells.ranges:
        if coord in rng:
            return rng.min_row, rng.min_col, rng.max_row, rng.max_col
    cell = sheet[coord]
    return cell.row, cell.column, cell.row, cell.column


def _value_right_of(sheet, coord: str) -> Optional[str]:
    """标签格右侧第一个非空格。标签本身常常是合并格（A20:D20 → 值在 E20）。"""
    min_row, _min_col, _max_row, max_col = _merged_span(sheet, coord)
    limit = min(sheet.max_column or max_col, max_col + 6)
    for col in range(max_col + 1, limit + 1):
        cell = sheet.cell(row=min_row, column=col)
        if str(cell.value or "").strip():
            return cell.coordinate
    return None


def _split_place(text: str) -> Tuple[str, str]:
    """A3 = 「排班时间：… 值班地点：X」→ (前半原文, X)。"""
    if PLACE_LABEL in text:
        head, tail = text.split(PLACE_LABEL, 1)
        return head.rstrip(), tail.strip().lstrip("：:").strip()
    return text.strip(), ""


def scan_text_slots(sheet) -> Dict[str, str]:
    """找出可变文案格坐标。按内容识别，换模板也能用。"""
    slots: Dict[str, str] = {}
    for row in sheet.iter_rows():
        for cell in row:
            text = str(cell.value or "").strip()
            if not text or text in (MARKER, "A3"):
                continue
            if "title" not in slots and "值班表" in text and len(text) <= 60:
                slots["title"] = cell.coordinate
            elif "place" not in slots and PLACE_LABEL in text:
                slots["place"] = cell.coordinate
            elif "teacher_phone" not in slots and "负责老师" in text and "电话" in text:
                coord = _value_right_of(sheet, cell.coordinate)
                if coord:
                    slots["teacher_phone"] = coord
            elif "teacher_name" not in slots and "负责老师" in text:
                coord = _value_right_of(sheet, cell.coordinate)
                if coord:
                    slots["teacher_name"] = coord
            elif "note" not in slots and text.startswith("备注"):
                slots["note"] = cell.coordinate
    return slots


def read_template_texts(template_path: Path | str) -> Dict[str, str]:
    """读出模板里当前的文案值（地点只取「值班地点：」后面那段）。"""
    workbook = openpyxl.load_workbook(str(template_path))
    try:
        sheet = workbook.active
        slots = scan_text_slots(sheet)
        out: Dict[str, str] = {}
        for field, coord in slots.items():
            raw = str(sheet[coord].value or "")
            out[field] = _split_place(raw)[1] if field == "place" else raw.strip()
        return out
    finally:
        workbook.close()


def _apply_texts(sheet, texts: Mapping[str, str], week: Optional[int], today: date) -> None:
    """把在线改过的文案写进工作表；没填的字段保持模板原样。"""
    slots = scan_text_slots(sheet)
    for field in ("title", "teacher_name", "teacher_phone", "note"):
        value = (texts.get(field) or "").strip()
        coord = slots.get(field)
        if coord and value:
            sheet[coord] = value

    coord = slots.get("place")
    if coord:
        original = str(sheet[coord].value or "")
        _head, place_original = _split_place(original)
        place = (texts.get("place") or "").strip() or place_original
        if week is not None:
            stamp = f"{TIME_LABEL}：{today.year} 年 {today.month} 月 第 {week} 周"
            sheet[coord] = f"{stamp}   {PLACE_LABEL}：{place}" if place else stamp
        elif place:
            sheet[coord] = f"{PLACE_LABEL}：{place}"


def _compose_header(original: object, week: int, today: date) -> str:
    """A3 是「排班时间 + 值班地点」的合并格，只改时间，地点原样保留。"""
    text = str(original or "")
    stamp = f"{TIME_LABEL}：{today.year} 年 {today.month} 月 第 {week} 周"
    if PLACE_LABEL in text:
        tail = text.split(PLACE_LABEL, 1)[1]
        return f"{stamp}   {PLACE_LABEL}{tail}"
    return stamp
