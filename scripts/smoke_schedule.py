"""本地冒烟（第三阶段）：管理员一键排班 + Excel 下载校验。

用法：.venv\\Scripts\\python.exe scripts\\smoke_schedule.py [week]
前置：dev server 以 JIAOWU_MODE=stub 运行，超管/管理员账号已注册并通过。
"""

from __future__ import annotations

import re
import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import httpx2 as httpx
import openpyxl

from app.xlsx import read_layout

BASE = sys.argv[2] if len(sys.argv) > 2 else "http://127.0.0.1:8000"
WEEK = int(sys.argv[1]) if len(sys.argv) > 1 else 16

ADMIN_NO, ADMIN_PW = "2025000002", "demo-pass-2"
MEMBERS = [
    ("2025000003", "demo-pass-3"),
    ("2025000004", "demo-pass-4"),
    ("2025000005", "demo-pass-5"),
]
TEMPLATE = Path(__file__).resolve().parent.parent / "勤工助学值班表 20XX-20XX -X第X周.xlsx"
OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "out"

results: list[tuple[str, bool, str]] = []


def check(label: str, ok: bool, detail: str = "") -> None:
    results.append((label, ok, detail))
    print(f"{'PASS' if ok else 'FAIL'}  {label}" + (f"  [{detail}]" if detail else ""))


def csrf_of(client, path: str) -> str:
    page = client.get(path)
    match = re.search(r'name="csrf-token" content="([^"]+)"', page.text)
    assert match and match.group(1), f"{path} 没有 CSRF 令牌"
    return match.group(1)


def main() -> int:
    with httpx.Client(base_url=BASE, follow_redirects=False, timeout=20) as adm:
        r = adm.post("/login", data={"student_no": ADMIN_NO, "password": ADMIN_PW})
        if r.status_code != 303:
            check("管理员登录", False, str(r.status_code))
            return 1
        check("管理员登录", True)

        # 1) 把成员注册进本部门并审批
        for no, pw in MEMBERS:
            with httpx.Client(base_url=BASE, follow_redirects=False, timeout=20) as stu:
                reg = stu.post("/register", data={"student_no": no, "password": pw, "admin_key": ADMIN_NO})
            check(f"成员 {no} 注册", reg.status_code in (303, 409), str(reg.status_code))

        page = adm.get("/admin")
        pending_ids = re.findall(r'name="user_id" value="(\d+)"', page.text)
        csrf = csrf_of(adm, "/admin")
        approved = 0
        for uid in dict.fromkeys(pending_ids):
            res = adm.post("/admin/decide",
                           data={"user_id": uid, "decision": "approve", "csrf": csrf})
            approved += 1 if res.status_code == 303 else 0
        check("审批本部门成员", approved >= 1 or not pending_ids,
              f"{approved} 人" if pending_ids else "本轮无待审")

        # 2) 全员总览
        overview = adm.get(f"/admin/overview?week={WEEK}")
        check("全员总览可打开", overview.status_code == 200 and "全员总览" in overview.text,
              str(overview.status_code))
        check("总览含课表/请假标记", "有课：" in overview.text)

        # 3) 一键排班（第 16 周）
        run = adm.post("/admin/schedule/run",
                       data={"week": WEEK, "max_per_person": 2, "special_person_id": "",
                             "blacklist": "", "csrf": csrf_of(adm, f"/admin/schedule?week={WEEK}")})
        check("一键排班提交", run.status_code == 303, str(run.status_code))

        schedule = adm.get(f"/admin/schedule?week={WEEK}")
        planned = len(re.findall(r"<tr><td>周", schedule.text))
        check("排班结果页有内容", planned > 0, f"{planned} 个班次")

        # 4) 下载并校验 Excel
        download = adm.get(f"/admin/schedule/download?week={WEEK}")
        check("下载 Excel", download.status_code == 200
              and "spreadsheetml" in download.headers.get("content-type", ""),
              download.headers.get("content-type", ""))

        OUT_DIR.mkdir(parents=True, exist_ok=True)
        target = OUT_DIR / f"值班表-第{WEEK}周.xlsx"
        target.write_bytes(download.content)

        layout = read_layout(TEMPLATE)
        workbook = openpyxl.load_workbook(BytesIO(download.content))
        sheet = workbook.active
        filled = {}
        for (day, slot), cell in layout.name_cells.items():
            value = sheet[cell].value
            if value:
                filled[(day, slot)] = str(value)
        sunday = [key for key in filled if key[0] == "周日"]
        # 模板里原本就写着上一轮的值班人；只要这些名字没出现在导出结果里，就说明旧数据被清干净了
        ORIGINAL_TEMPLATE_NAMES = {"演示管理员", "韩梅梅", "李雷", "王芳", "赵敏"}
        stale = sorted({
            str(sheet[cell].value)
            for (day, slot), cell in layout.name_cells.items()
            if str(sheet[cell].value or "").strip() in ORIGINAL_TEMPLATE_NAMES
        })
        workbook.close()

        check("Excel 里有排班", len(filled) > 0, f"{len(filled)} 格")
        check("Excel 不含周日", not sunday, str(sunday))
        check("旧模板姓名已清空", not stale, str(stale))
        header = str(openpyxl.load_workbook(target).active["A3"].value)
        check("表头写上周次", f"第 {WEEK} 周" in header, header[:40])
        check("表头保留值班地点", "值班地点" in header and "A2、A3-4楼" in header)

        print("\n本周排班预览：")
        for (day, slot), name in sorted(filled.items()):
            print(f"  {day} {slot:>6}  {name}")
        print(f"\n文件已保存：{target}")

        # 5) 指定周六固定人员后重排：剩余空格应当被他填满
        member_id = re.search(
            rf'<option value="(\d+)"[^>]*>\s*同学3403', adm.get("/admin/schedule").text
        )
        if member_id is None:
            options = re.findall(r'<option value="(\d+)"[^>]*>\s*([^<]+)</option>',
                                 adm.get("/admin/schedule").text)
            special_id = options[0][0] if options else ""
            special_name = options[0][1].strip() if options else ""
        else:
            special_id = member_id.group(1)
            special_name = "同学3403"

        if special_id:
            again = adm.post("/admin/schedule/run",
                             data={"week": WEEK, "max_per_person": 2,
                                   "special_person_id": special_id, "blacklist": "",
                                   "csrf": csrf_of(adm, f"/admin/schedule?week={WEEK}")})
            check("指定固定人员后重排", again.status_code == 303, str(again.status_code))

            retry = adm.get(f"/admin/schedule/download?week={WEEK}")
            workbook = openpyxl.load_workbook(BytesIO(retry.content))
            sheet = workbook.active
            filled2 = {}
            for (day, slot), cell in layout.name_cells.items():
                if sheet[cell].value:
                    filled2[(day, slot)] = str(sheet[cell].value)
            workbook.close()
            saturday = {name for (day, _slot), name in filled2.items() if day == "周六"}
            check("指定人员填满剩余格", len(filled2) >= len(filled),
                  f"{len(filled)} → {len(filled2)}")
            check("周六只排指定人员", saturday <= {special_name}, str(saturday))
            retry_target = OUT_DIR / f"值班表-第{WEEK}周-含固定人员.xlsx"
            retry_target.write_bytes(retry.content)
            print(f"含固定人员的完整表格：{retry_target}")

    failed = [r for r in results if not r[1]]
    print(f"{len(results) - len(failed)}/{len(results)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
